#!/usr/bin/env python3
import openpyxl

for filename in ['All_Controls_Catalog_26091416.xlsx', 'Unique_Controls_Catalog_26091416.xlsx']:
    file = rf'C:\PySC\TAP\Output\Processed\Normalized\{filename}'
    try:
        wb = openpyxl.load_workbook(file)
        print(f"\n{filename}:")
        for sheet in wb.sheetnames[:5]:  # Show first 5 sheets
            print(f"  - {sheet}")
        if len(wb.sheetnames) > 5:
            print(f"  ... and {len(wb.sheetnames) - 5} more")
    except Exception as e:
        print(f"Error reading {filename}: {e}")
